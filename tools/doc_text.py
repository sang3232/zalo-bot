# -*- coding: utf-8 -*-
"""Trích chữ từ file tài liệu cho bot Zalo (C:\\zalobot\\src\\docs.js gọi).

CHỈ ĐỌC dữ liệu bên trong file — không mở bằng Word/Excel, không chạy macro.
Dùng:  python doc_text.py <file> [thư_mục_ảnh]
In ra MỘT dòng JSON (ASCII): {"ok": true, "text": "...", "pages": 3, "images": [...], "note": ""}
PDF quét (trang không có lớp chữ) thì lưu ảnh nhúng vào thư_mục_ảnh để bot gửi OCR.
"""
import datetime
import json
import os
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

MAX_CHARS = 200_000
MAX_UNZIPPED = 150 * 1024 * 1024    # chặn file nén "bom" (docx/xlsx/pptx đều là zip)
MAX_PDF_PAGES = 200
MAX_IMAGES = 10
MIN_CHARS_PER_PAGE = 30             # ít chữ hơn thế coi là trang quét
MIN_IMAGE_PIXELS = 200_000          # bỏ logo, biểu tượng nhỏ (dải chữ 1200x260 vẫn giữ)


def check_zip(path):
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
            raise ValueError("nội dung giải nén quá lớn")


def fmt_cell(v):
    # Ngày trong Excel là số/datetime — đổi ra dd/mm/yyyy để bot dò được mốc
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y %H:%M") if (v.hour or v.minute) else v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_pdf(path, img_dir):
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("PDF có mật khẩu")
    total = len(reader.pages)
    parts, images, scanned = [], [], 0
    for i in range(min(total, MAX_PDF_PAGES)):
        page = reader.pages[i]
        text = (page.extract_text() or "").strip()
        parts.append(f"[Trang {i + 1}]\n{text}")
        if len(text) >= MIN_CHARS_PER_PAGE or not img_dir:
            continue
        scanned += 1
        try:
            for img in page.images:
                if len(images) >= MAX_IMAGES:
                    break
                pil = img.image
                if pil.size[0] * pil.size[1] < MIN_IMAGE_PIXELS:
                    continue
                out = os.path.join(img_dir, f"trang{i + 1}_{len(images)}.png")
                pil.convert("RGB").save(out, "PNG")
                images.append(out)
        except Exception:
            pass
    note = f"{scanned}/{min(total, MAX_PDF_PAGES)} trang không có lớp chữ" if scanned else ""
    return "\n\n".join(parts), total, images, note


def read_docx(path):
    check_zip(path)
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    # Đi theo đúng thứ tự trong văn bản: đoạn văn và bảng xen nhau
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, d).text.strip()
            if t:
                out.append(t)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = []
                for c in row.cells:
                    s = c.text.strip()
                    if s and (not cells or cells[-1] != s):   # ô gộp bị lặp lại
                        cells.append(s)
                if cells:
                    out.append(" | ".join(cells))
    return "\n".join(out)


def read_xlsx(path):
    check_zip(path)
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out, size = [], 0
    try:
        for ws in wb.worksheets:
            out.append(f"[Trang tính: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [fmt_cell(v) for v in row if v is not None and str(v).strip() != ""]
                if cells:
                    line = " | ".join(cells)
                    out.append(line)
                    size += len(line)
                if size > MAX_CHARS:
                    break
    finally:
        wb.close()
    return "\n".join(out)


def read_pptx(path):
    check_zip(path)
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        slides = [n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]
        slides.sort(key=lambda n: int(re.search(r"slide(\d+)\.xml$", n).group(1)))
        out = []
        for i, name in enumerate(slides, 1):
            root = ET.fromstring(z.read(name))
            texts = [t.text for t in root.iter(ns + "t") if t.text]
            out.append(f"[Trang chiếu {i}]\n" + " ".join(texts))
    return "\n\n".join(out)


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "thiếu đường dẫn file"}))
        return
    path = sys.argv[1]
    img_dir = sys.argv[2] if len(sys.argv) > 2 else ""
    ext = os.path.splitext(path)[1].lower()
    try:
        images, pages, note = [], None, ""
        if ext == ".pdf":
            text, pages, images, note = read_pdf(path, img_dir)
        elif ext == ".docx":
            text = read_docx(path)
        elif ext == ".xlsx":
            text = read_xlsx(path)
        elif ext == ".pptx":
            text = read_pptx(path)
        else:
            raise ValueError(f"không hỗ trợ đuôi {ext}")
        print(json.dumps({"ok": True, "text": text[:MAX_CHARS], "pages": pages,
                          "images": images, "note": note}, ensure_ascii=True))
    except Exception as e:
        print(json.dumps({"ok": False, "error": f"{type(e).__name__}: {e}"}, ensure_ascii=True))


if __name__ == "__main__":
    main()
