# -*- coding: utf-8 -*-
"""Tạo file mẫu cho test/docs.test.js: python tao_mau.py <thư_mục>"""
import datetime
import os
import sys
import zipfile

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

# DOCX: đoạn văn + bảng
import docx
d = docx.Document()
d.add_heading("THÔNG BÁO LỊCH HỌP PHỤ HUYNH", 1)
d.add_paragraph("Thời gian: 8h30 sáng Thứ Bảy, ngày 10/10/2026")
t = d.add_table(rows=2, cols=3)
for i, s in enumerate(["Môn", "Ngày thi", "Giờ"]):
    t.rows[0].cells[i].text = s
for i, s in enumerate(["Toán", "12/10/2026", "07:30"]):
    t.rows[1].cells[i].text = s
d.save(os.path.join(out, "lich_hop.docx"))

d2 = docx.Document()
d2.add_paragraph("Biên bản buổi vui chơi tập thể, mọi người rất vui.")
d2.save(os.path.join(out, "khong_lich.docx"))

# XLSX: ô ngày giờ kiểu datetime
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Lịch thi"
ws.append(["Môn", "Thời gian", "Phòng"])
ws.append(["Toán", datetime.datetime(2026, 10, 12, 7, 30), "P.302"])
ws.append(["Văn", datetime.date(2026, 10, 13), "P.301"])
wb.save(os.path.join(out, "lich_thi.xlsx"))

# PPTX tối thiểu: chỉ cần các slide XML
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
def slide(text):
    return (f'<?xml version="1.0" encoding="UTF-8"?><p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            f'xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
with zipfile.ZipFile(os.path.join(out, "ke_hoach.pptx"), "w") as z:
    z.writestr("ppt/slides/slide2.xml", slide("Hạn nộp báo cáo: 20/10/2026"))
    z.writestr("ppt/slides/slide1.xml", slide("Kế hoạch tháng 10"))

# PDF có lớp chữ (viết tay, font Helvetica, chữ không dấu)
def make_pdf(path, text):
    content = f"BT /F1 18 Tf 50 750 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    buf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(buf))
        buf += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(buf)
    buf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        buf += b"%010d 00000 n \n" % off
    buf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    with open(path, "wb") as f:
        f.write(buf)
make_pdf(os.path.join(out, "thong_bao.pdf"), "Hop to dan pho ngay 14/10/2026 luc 19:30")

# Ảnh chữ + PDF quét (chỉ có ảnh, không có lớp chữ) — cho đường OCR
from PIL import Image, ImageDraw, ImageFont
font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 40)
img = Image.new("RGB", (1200, 260), "white")
dr = ImageDraw.Draw(img)
dr.text((40, 40), "Lịch khám sức khoẻ: 9h00 ngày 15/10/2026", fill="black", font=font)
dr.text((40, 130), "Địa điểm: Trạm y tế phường", fill="black", font=font)
img.save(os.path.join(out, "anh_lich.png"))
img.save(os.path.join(out, "scan.pdf"), "PDF", resolution=150)

with open(os.path.join(out, "ghi_chu.txt"), "w", encoding="utf-8") as f:
    f.write("Nộp báo cáo quý trước 20/10/2026")
with open(os.path.join(out, "virus.exe"), "wb") as f:
    f.write(b"MZ-khong-phai-chuong-trinh-that")
print("ok")
